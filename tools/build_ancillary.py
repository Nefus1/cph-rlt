"""Notary acknowledgment, Certification of Trust, General Assignment,
Pour-Over Will, and Plan Summary memo templates."""

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt

from docbuilder import Builder
from common import X, x, DATED, DATE, COUNTY, TRUSTID


# ---------------------------------------------------------------- notary
def build_notary(path):
    """California all-purpose acknowledgment (Civil Code section 1189)."""
    raw = Builder(numbered=False)
    b = X(raw)
    raw.page_break()
    b.para('CALIFORNIA ALL-PURPOSE ACKNOWLEDGMENT', 'Center Bold', after=12)
    t = raw.doc.add_table(rows=1, cols=1)
    t.style = raw.doc.styles['Table Grid']
    cell = t.rows[0].cells[0]
    cell.paragraphs[0].style = raw.doc.styles['Small']
    raw._runs(cell.paragraphs[0],
              'A notary public or other officer completing this certificate verifies only the '
              'identity of the individual who signed the document to which this certificate is '
              'attached, and not the truthfulness, accuracy, or validity of that document.')
    cell.width = Inches(4.0)
    raw.doc.add_paragraph(style='Normal')
    b.para_left('State of California<tab><tab><tab>)', after=0)
    b.para_left('County of ' + COUNTY + '<tab>)', after=12)
    b.para_left('On ____________________ before me, ______________________________________________ '
            '(insert name and title of the officer), personally appeared '
            '________________________________________________________________, who proved to me '
            'on the basis of satisfactory evidence to be the person(s) whose name(s) is/are '
            'subscribed to the within instrument and acknowledged to me that he/she/they '
            'executed the same in his/her/their authorized capacity(ies), and that by '
            'his/her/their signature(s) on the instrument the person(s), or the entity upon '
            'behalf of which the person(s) acted, executed the instrument.')
    b.para_left('I certify under PENALTY OF PERJURY under the laws of the State of California that the '
            'foregoing paragraph is true and correct.')
    b.para_left('WITNESS my hand and official seal.', after=30)
    b.para_left('Signature _______________________________________<tab><tab>(Seal)')
    raw.save(path)


# ---------------------------------------------------------------- certification
def build_certification(path):
    raw = Builder(numbered=False)
    b = X(raw)
    raw.footer('Certification of Trust | {{ trust_title }}')
    b.title('Certification of Trust')
    b.subtitle('California Probate Code Section 18100.5')
    b.center('<b>{{ trust_title }}</b><br>' + DATED, after=18)
    b.plain('The undersigned, {{ s1_name | upper }} and {{ s2_name | upper }}, being all of the '
            'currently acting trustees of the trust described below, certify as follows:')

    def item(n, heading, text):
        b.para(n + '.<tab><b>' + heading + '.</b> ' + text, 'Hang')

    item('1', 'Existence of the Trust',
         'The trust known as {{ trust_title }} exists and was executed on '
         '{% if is_restatement %}{{ original_dated }}, and was amended and restated in its '
         'entirety on @DATE@{% else %}@DATE@{% endif %}.')
    item('2', 'Settlors',
         'The settlors (@Ts@) of the trust are {{ s1_name }} and {{ s2_name }}.')
    item('3', 'Currently Acting Trustees',
         'The currently acting trustees are {{ s1_name }} and {{ s2_name }}, whose address is '
         '{{ residence.on_one_line() }}.')
    item('4', 'Powers of the Trustee',
         'The trustees have full power to acquire, sell, assign, convey, exchange, encumber, '
         'lease, borrow against, invest, and otherwise deal with any interest in real or personal '
         'property held in the trust; to open, maintain, and close accounts of every kind with '
         'banks, brokers, and other financial institutions; to sign checks and withdraw funds; and '
         'to exercise all powers granted by the trust instrument and by California Probate Code '
         'Sections 16200 through 16249.')
    item('5', 'Power to Revoke',
         'The trust is revocable. While both settlors are living, either settlor may revoke the '
         'trust as to community property, and each settlor may revoke the trust as to that '
         'settlor\'s separate property. After the death of a settlor, the surviving settlor may '
         'revoke the survivor\'s trust.')
    item('6', 'Signature Authority',
         '{% if joint_trustee_authority == "either" %}When more than one trustee is acting, any one '
         'trustee acting alone may sign and exercise any power of the trustee.{% else %}When more '
         'than one trustee is acting, the signatures of all acting trustees are required to '
         'exercise the powers of the trustee.{% endif %}')
    item('7', 'Taxpayer Identification Number',
         'The trust\'s taxpayer identification number is '
         '{% if tin_mode == "enter" %}{{ trust_tin }}{% else %}<ph>______________________</ph>{% endif %}.')
    item('8', 'Title to Trust Property',
         'Title to trust property shall be taken as: "{{ s1_name }} and {{ s2_name }}, Trustees of '
         '@TRUSTID@."')
    b.ctl('{%p if realprop_assets | length %}')
    item('9', 'Real Property',
         'The trust holds an interest in the following real property:')
    b.ctl('{%p for a in realprop_assets %}')
    b.para('{{ a.description | trim }}{% if a.apn %} (APN {{ a.apn }}){% endif %}: '
           '{% if a.legal_description %}{{ a.legal_description | trim | replace("\\n", " ") }}'
           '{% else %}<ph>[Insert legal description from the vesting deed]</ph>{% endif %}',
           'Body Indent')
    b.ctl('{%p endfor %}')
    b.ctl('{%p endif %}')
    item('{{ 10 if realprop_assets | length else 9 }}', 'No Revocation or Amendment',
         'The trust has not been revoked, modified, or amended in any manner that would cause the '
         'representations in this certification to be incorrect.')
    b.plain('We declare under penalty of perjury under the laws of the State of California that '
            'the foregoing is true and correct.', before=6, keep=True)
    b.plain('Executed on @DATE@, at @COUNTY@ County, California.', keep=True)
    b.sig('<b>{{ s1_name | upper }}</b>', 'Trustee')
    b.sig('<b>{{ s2_name | upper }}</b>', 'Trustee')
    b.ctl("{{p include_docx_template('notary_ack_ca.docx') }}")
    raw.save(path)


# ---------------------------------------------------------------- assignment
def build_assignment(path):
    raw = Builder(numbered=False)
    b = X(raw)
    raw.footer('General Assignment | {{ trust_title }}')
    b.title('General Assignment of Assets')
    b.subtitle('To {{ trust_title }}, ' + DATED)
    b.plain('For valuable consideration, the receipt of which is acknowledged, {{ s1_name | upper }} '
            'and {{ s2_name | upper }} (the "Assignors") assign, transfer, and convey to '
            '{{ s1_name }} and {{ s2_name }}, Trustees of @TRUSTID@ (the "Trust"), all of the '
            'Assignors\' right, title, and interest in the following property, wherever located, '
            'whether now owned or later acquired:')

    def item(n, text):
        b.para(n + '.<tab>' + text, 'Hang')

    item('1', 'All tangible personal property, including household furniture and furnishings, '
              'jewelry, clothing, books, art, collections, tools, and personal effects;')
    item('2', 'All deposit, savings, money market, brokerage, and investment accounts, '
              'certificates of deposit, stocks, bonds, notes, and other securities, and all cash;')
    item('3', 'All interests in corporations, limited liability companies, partnerships, and other '
              'business entities, subject to any restrictions on transfer in their governing '
              'documents;')
    item('4', 'All copyrights, royalties, intellectual property, digital assets, and claims, '
              'refunds, and causes of action;')
    item('5', 'All real property described on the Schedules to the Trust, and every other interest '
              'in real property that the Assignors intend to be held in the Trust; and')
    item('6', 'All other property of every kind that is not effectively transferred to the Trust '
              'by another instrument.')
    b.plain('<b>Exclusions.</b> This assignment does not include (a) interests in qualified '
            'retirement plans, individual retirement accounts, tax-sheltered annuities, health '
            'savings accounts, and education savings accounts; (b) life insurance policies and '
            'annuities; or (c) any property that would lose a tax benefit or create a default '
            'under its governing documents by reason of this assignment. Those assets pass by '
            'their own beneficiary designations.', before=6)
    b.plain('<b>Character of Property.</b> Community property assigned by this instrument remains '
            'community property, and the separate property of each Assignor remains that '
            'Assignor\'s separate property, as provided in the Trust and Family Code Section 761.')
    b.plain('<b>Intent.</b> The Assignors intend this instrument to be an effective declaration '
            'that the property described above is held in the Trust, to be relied on in any '
            'proceeding under Probate Code Section 850. The Assignors agree to sign any further '
            'documents needed to complete the transfer of title to the Trustees.')
    b.plain('Signed on @DATE@, at @COUNTY@ County, California.', keep=True, before=6)
    b.para('<b>ASSIGNORS:</b>', 'Plain', keep=True)
    b.sig('<b>{{ s1_name | upper }}</b>', 'Assignor')
    b.sig('<b>{{ s2_name | upper }}</b>', 'Assignor')
    b.para('<b>ACCEPTED BY THE TRUSTEES:</b>', 'Plain', keep=True, before=18)
    b.sig('<b>{{ s1_name | upper }}</b>', 'Trustee')
    b.sig('<b>{{ s2_name | upper }}</b>', 'Trustee')
    b.ctl("{{p include_docx_template('notary_ack_ca.docx') }}")
    raw.save(path)


# ---------------------------------------------------------------- pour-over will
TESTATOR = '(spouse1 if will_for | string == "1" else spouse2)'


def build_will(path):
    raw = Builder(numbered=True)
    b = X(raw)
    raw.footer('Last Will of {{ ' + TESTATOR + '.name.full(middle="full") }}',
               initials='          Initials: ________')
    b.ctl('{%p set testator = ' + TESTATOR + ' %}')
    b.ctl('{%p set spouse = (spouse2 if will_for | string == "1" else spouse1) %}')
    b.title('Last Will of {{ testator.name.full(middle="full") | upper }}')
    b.subtitle('Pour-Over Will')
    b.plain('I, <b>{{ testator.name.full(middle="full") | upper }}</b>, a resident of '
            '{{ county_name }} County, California, declare that this is my Will. I revoke all '
            'prior wills and codicils.', before=12)

    b.article('Family')
    b.section('Spouse',
              'I am {% if union_type == "rdp" %}the registered domestic partner of{% else %}married '
              'to{% endif %} {{ spouse.name.full(middle="full") }}, referred to in this Will as my '
              'spouse.')
    b.ctl('{%p if children | length %}')
    b.section('Children',
              '{% if children | length == 1 %}The following person is{% else %}The following '
              'persons are{% endif %} named as {% if children | length == 1 %}a child{% else %}'
              'children{% endif %} in our trust:')
    b.ctl('{%p for c in children %}')
    b.sub('{{ c.name.full(middle="full") }}{% if (c.parentage == "s1" and will_for | string == "2") '
          'or (c.parentage == "s2" and will_for | string == "1") %}, my spouse\'s child'
          '{% if stepchildren_equal %}, whom I treat as my own child for all purposes of this '
          'Will{% endif %}{% endif %}{% if not c.is_living %} (deceased){% endif %}'
          '{% if loop.last %}.{% else %};{% endif %}')
    b.ctl('{%p endfor %}')
    b.ctl('{%p else %}')
    b.section('Children', 'I have no children.')
    b.ctl('{%p endif %}')
    b.section('Intentional Omission',
              'Except as provided in this Will and in the trust described below, I have '
              'intentionally, and with full knowledge, omitted to provide for my heirs.'
              '{% if omitted_people | length %} Without limitation, I make no provision for '
              '{{ join_list(omitted_people, "or") }}.{% endif %}')

    b.article('Gift to Trust')
    b.section('Pour-Over Gift',
              'I give my entire estate to the trustee then serving under {{ trust_title }}, '
              '@DATED@, which my spouse and I signed as @Ts@ before signing this Will, to be added '
              'to and administered as part of that trust according to its terms, including any '
              'amendments made before my death. This gift is made under Probate Code Section 6300 '
              'and shall not be invalid because the trust is amendable or revocable.')
    b.section('Incorporation by Reference',
              'If the gift to the trust fails for any reason, I give my estate to the trustee named '
              'in that trust, to be held, administered, and distributed under the terms of the '
              'trust as they exist on the date of this Will, which I incorporate by reference.')
    b.section('Powers of Appointment',
              'I do not exercise any power of appointment by this Will, except as I may expressly '
              'state in a later codicil.')

    b.article('Executor')
    b.section('Nomination',
              'I nominate my spouse as executor of this Will. If my spouse fails to qualify or '
              'ceases to act, I nominate the following, in the order listed:')
    b.ctl('{%p for e in will_executors %}')
    b.sub('{{ e.described() }}{% if loop.last %}.{% else %};{% endif %}')
    b.ctl('{%p endfor %}')
    b.section('Bond and Independent Administration',
              'No bond shall be required of any executor. My executor shall have full authority to '
              'administer my estate under the Independent Administration of Estates Act (Probate '
              'Code Section 10400 and following).')
    b.section('Tax Elections',
              'My executor may make any tax election, including the election to transfer my unused '
              'exclusion amount to my spouse under Section 2010(c)(5)(A) of the Internal Revenue '
              'Code, without adjustment among beneficiaries.')

    b.ctl('{%p if minor_children_exist and nominate_guardians %}')
    b.article('Guardian')
    b.section('Nomination of Guardian',
              'If my spouse does not survive me or is unable to act, I nominate the following, in '
              'the order listed, as guardian of the person and estate of each minor child of mine, '
              'under Probate Code Section 1500 and following:')
    b.ctl('{%p for g in guardians %}')
    b.sub('{{ g.described() }}{% if g.co_guardian %}, to serve jointly with {{ g.co_guardian }}'
          '{% endif %}{% if loop.last %}.{% else %};{% endif %}')
    b.ctl('{%p endfor %}')
    b.section('No Bond', 'No bond shall be required of any guardian nominated in this Will.')
    b.ctl('{%p endif %}')

    b.article('General Provisions')
    b.section('Debts and Taxes',
              'My executor may request that the trustee of the trust pay my debts, expenses of '
              'administration, and any estate taxes, as provided in the trust.')
    b.ctl('{%p if include_no_contest %}')
    b.section('No-Contest Clause',
              'Any beneficiary who, without probable cause, brings a direct contest (as defined in '
              'Probate Code Section 21310) of this Will, the trust, or any other protected '
              'instrument, or who files a creditor\'s claim or a pleading challenging a transfer on '
              'the ground that the property was not mine, shall take nothing under this Will. This '
              'Section shall be enforced only to the extent permitted by Probate Code Sections '
              '21310 through 21315.')
    b.ctl('{%p endif %}')
    b.section('Governing Law', 'This Will shall be construed under California law.')

    b.para_left('I subscribe my name to this Will on ________________________, 20____, at '
            '________________________, California, and declare under penalty of perjury under the '
            'laws of the State of California that I sign and execute this instrument as my Will, '
            'that I sign it willingly, that I execute it as my free and voluntary act, and that I '
            'am eighteen years of age or older, of sound mind, and under no constraint or undue '
            'influence.', before=12, keep=True)
    b.sig('<b>{{ testator.name.full(middle="full") | upper }}</b>', 'Testator')

    b.para('ATTESTATION AND DECLARATION OF WITNESSES', 'Center Bold', before=24)
    b.plain('On the date written below, {{ testator.name.full(middle="full") }} (the "Testator") '
            'declared to us, the undersigned, that this instrument, consisting of '
            '<field:NUMPAGES> pages including the page signed by us as witnesses, is the '
            'Testator\'s Will and asked us to act as witnesses to it. The Testator then signed '
            'this Will in our presence, all of us being present at the same time. We now, at the '
            'Testator\'s request, in the Testator\'s presence, and in the presence of each other, '
            'sign our names as witnesses. Each of us is now eighteen years of age or older and a '
            'competent witness, and resides at the address written after our name. We understand '
            'that this instrument is the Testator\'s Will. To the best of our knowledge, the '
            'Testator is eighteen years of age or older, is of sound mind, and is under no '
            'constraint or undue influence of any kind. Neither of us is named as a beneficiary '
            'in this Will.')
    b.plain('We declare under penalty of perjury under the laws of the State of California that '
            'the foregoing is true and correct.', keep=True)
    b.para_left('Executed on ________________________, 20____, at ________________________, '
                'California.', keep=True)
    for n in ('1', '2'):
        b.para_left('Witness ' + n + ' signature: _______________________________________________',
                    before=24, after=6, keep=True)
        b.para_left('Printed name: _______________________________________________________',
                    after=6, keep=True)
        b.para_left('Residence address: _________________________________________________________',
                    after=0, keep=(n == '1'))
    raw.save(path)


# ---------------------------------------------------------------- summary memo
def build_summary(path):
    raw = Builder(numbered=False, space_after=6)
    b = X(raw)

    def table(headers, rows, widths=None, loop=None):
        headers = [x(h) for h in headers] if headers else headers
        rows = [[x(c) for c in row] for row in rows]
        if loop:
            loop = (loop[0], [x(c) for c in loop[1]])
        return raw.table(headers, rows, widths=widths, loop=loop)
    raw.footer('Plan Summary | {{ trust_title }} | Privileged and Confidential')
    b.title('Plan Summary and Funding Memo')
    b.subtitle('Privileged and confidential | Attorney work product | Draft for attorney review')
    table(None, [
        ['<b>Clients</b>', '{{ s1_name }} and {{ s2_name }}'],
        ['<b>Trust</b>', '@TRUSTID@'],
        ['<b>Residence</b>', '{{ residence.on_one_line() }} ({{ county_name }} County)'],
        ['<b>Prepared by</b>', '{% if preparer_attorney %}{{ preparer_attorney }}'
                               '{% if preparer_bar %} (SBN {{ preparer_bar }}){% endif %}{% else %}'
                               '<ph>[Attorney name]</ph>{% endif %}, {% if preparer_firm %}'
                               '{{ preparer_firm }}{% else %}<ph>[Firm name]</ph>{% endif %}'
                               '{% if preparer_phone %} | {{ preparer_phone }}{% endif %}'
                               '{% if preparer_email %} | {{ preparer_email }}{% endif %}'],
        ['<b>Memo date</b>', '{{ format_date(today(), format="MMMM d, yyyy") }}'],
    ], widths=[1.5, 5.0])

    b.para('1. The Plan at a Glance', 'Memo Heading')
    table(['Decision', 'Choice'], [
        ['Trust type', 'Joint revocable living trust{% if is_restatement %} (restatement of '
                       '{{ original_trust_name }} dated {{ original_dated }}){% endif %}'],
        ['Plan at first death', '{{ plan_label }}'],
        ['Survivor\'s control', '{% if plan_type == "survivor" %}Survivor controls the entire trust '
                                'and may change beneficiaries.{% elif plan_type == "disclaimer" %}'
                                'Survivor controls everything not disclaimed within nine months.'
                                '{% else %}Survivor controls only the Survivor\'s Trust; the Bypass '
                                'Trust is locked in at the first death.{% endif %}'],
        ['{{ subtrust_name or "Sub-trust" }} terms', '{% if plan_type in ("ab", "disclaimer") %}'
                                '{{ "All net income" if subtrust_income == "mandatory" else "Income for HEMS" }}'
                                ', principal for HEMS{% if plan_type == "ab" and bypass_five_five %}, '
                                '5-and-5 power{% endif %}{% if plan_type == "ab" and bypass_lpoa %}, '
                                'limited power of appointment{% endif %}{% else %}Not applicable'
                                '{% endif %}'],
        ['Specific gifts', '{% if specific_gifts | length %}{% for g in specific_gifts %}'
                           '{{ g.short() }}{% if not loop.last %}; {% endif %}{% endfor %}'
                           '{% else %}None{% endif %}'],
        ['Residue', '{% if residue_plan == "children" %}Children in equal shares{% else %}'
                    '{% for r in residuary %}{{ r.name.text }} ({{ "equal" if residuary_share_method == "equal" else pct(r.share) }})'
                    '{% if not loop.last %}; {% endif %}{% endfor %}{% endif %}, '
                    '{{ "per stirpes (Prob. Code 246)" if representation == "per_stirpes" else "per capita at each generation (Prob. Code 247)" }}'],
        ['Beneficiary shares', '{% if dist_method == "outright" %}Outright, held in trust until age '
                               '{{ young_age }}{% elif dist_method == "staged" %}Staged at ages '
                               '{% for s in stage_plan %}{{ s["age"] }}{% if not loop.last %}, {% endif %}'
                               '{% endfor %}{% else %}Lifetime protective trusts{% endif %}'],
        ['Supplemental needs', '{% if has_special_needs %}{{ special_needs_list | join(", ") }}'
                               '{% else %}None required (discretionary SNT clause included){% endif %}'],
        ['Backstop', '{% if ultimate_plan == "heirs_split" %}One-half to each spouse\'s heirs'
                     '{% elif ultimate_plan == "heirs" %}Heirs of the surviving spouse{% else %}'
                     '{{ ultimate_named | trim }}{% endif %}'],
        ['Options', 'No-contest clause: {{ "yes" if include_no_contest else "no" }}; '
                    'survivorship: {{ survivorship_days }} days; firearms: '
                    '{{ "yes" if has_firearms else "no" }}; CP agreement: '
                    '{{ "yes" if cp_presumption else "no" }}'],
    ], widths=[1.8, 4.7])

    b.para('2. Fiduciaries', 'Memo Heading')
    table(['Role', 'Who'], [
        ['Initial trustees', '{{ s1_name }} and {{ s2_name }} '
                             '({{ "either may act alone" if joint_trustee_authority == "either" else "both must sign" }})'],
        ['Successor trustees', '{% for t in successor_trustees %}{{ loop.index }}. {{ t.described() }}'
                               '{% if not loop.last %}<br>{% endif %}{% endfor %}'
                               '{% if successor_trustees | length > 1 %}<br><i>Serving '
                               '{{ "in order" if successor_mode == "sequential" else "together" }}</i>{% endif %}'],
        ['Executors', '{% if include_wills %}Surviving spouse, then {% for e in will_executors %}'
                      '{{ e.name.text }}{% if not loop.last %}, {% endif %}{% endfor %}{% else %}'
                      'No pour-over wills prepared{% endif %}'],
        ['Guardians', '{% if include_wills and minor_children_exist and nominate_guardians %}'
                      '{% for g in guardians %}{{ g.name.text }}{% if g.co_guardian %} and '
                      '{{ g.co_guardian }}{% endif %}{% if not loop.last %}, then {% endif %}'
                      '{% endfor %}{% elif minor_children_exist %}<ph>None nominated</ph>'
                      '{% else %}Not needed (no minor children){% endif %}'],
        ['Incapacity', '{{ incapacity_physicians }} physician certificate(s), or a court order'],
    ], widths=[1.8, 4.7])

    b.para('3. Placeholders and Blanks to Complete', 'Memo Heading')
    b.plain('Every item below is printed in the documents as a <ph>highlighted blank</ph>, or is '
            'a line completed at signing. Nothing else in the documents is a placeholder.')
    table(['Item', 'Where it appears', 'How to complete'], [],
              widths=[1.9, 2.4, 2.2],
              loop=('for p in placeholders', ['<b>{{ p["item"] }}</b>', '{{ p["where"] }}', '{{ p["fix"] }}']))

    b.para('4. Attorney Review Flags', 'Memo Heading')
    flags = [
        ('True', 'Joint representation: confirm a signed engagement letter with a joint '
                 'representation and conflict-waiver disclosure (Rule of Professional Conduct 1.7).'),
        ('plan_type == "survivor" and any_stepchildren',
         'Blended family with an All-to-Survivor plan: the surviving spouse can change the '
         'beneficiaries and disinherit the other spouse\'s children. Consider the Disclaimer or '
         'A/B plan.'),
        ('plan_type == "ab"',
         'Mandatory A/B: confirm the estate size justifies the split. The Bypass Trust needs an '
         'EIN and an annual Form 1041, and its assets get no second basis step-up at the '
         'survivor\'s death. Elect portability anyway.'),
        ('plan_type == "disclaimer"',
         'Disclaimer plan: at the first death, calendar the nine-month disclaimer deadline and '
         'instruct the survivor not to accept benefits from property that may be disclaimed.'),
        ('cp_presumption',
         'Community property agreement included: this is a transmutation (Fam. Code 852) that '
         'may benefit one spouse at the other\'s expense (Fam. Code 721 fiduciary duty). Document '
         'the advice given and the opportunity for independent counsel.'),
        ('has_special_needs',
         'Supplemental needs trusts required for: {{ special_needs_list | join(", ") }}. Make sure '
         'beneficiary designations (life insurance, retirement accounts) name the trust, not the '
         'beneficiary.'),
        ('residue_plan == "custom" and residuary | selectattr("is_charity") | list | length',
         'Charitable beneficiaries: confirm each charity\'s exact legal name and EIN.'),
        ('specific_gifts | selectattr("is_charity") | list | length',
         'Charitable specific gifts: confirm each charity\'s exact legal name and EIN.'),
        ('minor_children_exist and not (include_wills and nominate_guardians)',
         'Minor children but no guardian nomination: a guardian can be nominated only in a will '
         'or other signed writing (Prob. Code 1500-1502).'),
        ('designation_assets | length',
         'Retirement accounts or life insurance listed: review the beneficiary designations and '
         'SECURE Act payout rules before naming the trust.'),
        ('realprop_assets | length',
         'Real property: prepare and record a trust transfer deed for each parcel, with a PCOR. '
         'Out-of-state real property needs its own deed to avoid ancillary probate.'),
        ('True',
         'Property tax: transfers to the trust are excluded from reassessment (R&T Code 62(d)). '
         'A later transfer of the home to children may qualify for the limited Proposition 19 '
         'parent-child exclusion (R&T Code 63.2); the claim must be filed after the death.'),
        ('has_firearms',
         'Firearms: advise the successor trustees about California transfer rules; consider a '
         'separate firearms list.'),
        ('is_restatement',
         'Restatement: confirm that the original trust\'s amendment provisions allow a restatement '
         'by the method used, and that both spouses are the original trustors.'),
        ('True',
         'Ancillary documents not prepared by this questionnaire: durable powers of attorney, '
         'advance health care directives, HIPAA releases, and trust transfer deeds.'),
    ]
    for cond, text in flags:
        b.ctl('{%p if ' + cond + ' %}')
        b.para('•<tab>' + text, 'Bullet Plain')
        b.ctl('{%p endif %}')

    b.para('5. Funding Checklist', 'Memo Heading')
    b.ctl('{%p if assets | length %}')
    table(['Asset', 'Character', 'Action'], [], widths=[2.2, 1.4, 2.9],
              loop=('for a in assets', ['{{ a.description }}{% if a.asset_type == "realprop" and a.apn %} (APN {{ a.apn }}){% endif %}',
                                        '{{ a.character_label(s1_name, s2_name) }}',
                                        '{{ a.funding_step(trust_title) }}']))
    b.ctl('{%p else %}')
    b.plain('No assets were listed. Gather the deed for each parcel, recent statements for every '
            'account, and the beneficiary designation for every retirement account and insurance '
            'policy, then retitle each asset as described below.')
    b.ctl('{%p endif %}')
    for txt in [
        'Real property: trust transfer deed, PCOR (BOE-502-A), documentary transfer tax exemption '
        'noted (R&T Code 11930).',
        'Bank and brokerage accounts: retitle to "{{ s1_name }} and {{ s2_name }}, Trustees of '
        '@TRUSTID@," using the Certification of Trust.',
        'Retirement accounts and life insurance: do not retitle; update beneficiary designations.',
        'Vehicles and household goods: covered by the General Assignment; DMV retitling is optional.',
    ]:
        b.para('•<tab>' + txt, 'Bullet Plain')

    b.para('6. Signing Instructions', 'Memo Heading')
    for txt in [
        'Trust Agreement: both spouses sign and initial each page; sign before a notary '
        '(acknowledgment attached).',
        'Certification of Trust and General Assignment: both spouses sign before a notary.',
        '{% if include_wills %}Pour-over wills: each spouse signs in the presence of two adult '
        'witnesses who are present at the same time (Prob. Code 6110). Use witnesses who are not '
        'beneficiaries (Prob. Code 6112). Wills are not notarized.{% else %}No wills were '
        'prepared; confirm the clients have current pour-over wills.{% endif %}',
        'Keep the original trust and wills in a safe place known to the successor trustee; give '
        'copies of the Certification of Trust to financial institutions.',
    ]:
        b.para('•<tab>' + txt, 'Bullet Plain')

    b.para('7. After the First Death', 'Memo Heading')
    for txt in [
        'Serve the trustee\'s notification under Prob. Code 16061.7 within 60 days (only for '
        'any part of the trust that becomes irrevocable).',
        'File a Change in Ownership Statement (Death of Real Property Owner, BOE-502-D) within '
        '150 days (R&T Code 480).',
        'Consider the portability election on Form 706: due nine months after death (fifteen '
        'with extension), or up to five years for estates not otherwise required to file '
        '(Rev. Proc. 2022-32).',
        '{% if plan_type == "disclaimer" %}Decide on any disclaimer within nine months, before '
        'the survivor accepts benefits.{% elif plan_type == "ab" %}Fund the Bypass Trust, obtain '
        'its EIN, and file Form 1041 each year.{% else %}Retitle the deceased spouse\'s interest '
        'if needed and obtain a basis step-up appraisal.{% endif %}',
    ]:
        b.para('•<tab>' + txt, 'Bullet Plain')
    raw.save(path)
